"""
core/os/cognition/escl/three_brain_orchestrator.py
Three-Brain Orchestration Engine (Cognitive Brain + Capability Brain + Execution Brain).

Enforces the North Star Equation:
EXECUTABLE MISSION = Semantic Completeness * Capability Awareness * Contract Compatibility * Authorization * Verification
"""

from __future__ import annotations
import os
from typing import Any, Dict, Optional
from core.os.cognition.escl.canonical_mission import CanonicalMissionSpec
from core.os.cognition.escl.capability_brain import capability_brain, CapabilityAssessment
from core.os.cognition.escl.schema_adapter import schema_adapter
from core.os.cognition.escl.artifact_semantic_verifier import artifact_semantic_verifier
from core.os.cognition.fast_governor import fast_governor
from core.os.cognition.fast_schemas import ActionIntent, TaskDomain
from intelligence.skills.BUSINESS.OFFICE_SUITE_MASTER.smart_office_adapter import smart_office_adapter


class ThreeBrainOrchestrator:
    """End-to-end coordinator for the 3 Sub-Brains."""

    def process_mission(
        self,
        goal_text: str,
        mission_id: str = "m_default",
        output_dir: Optional[str] = None,
    ) -> Dict[str, Any]:
        # ── BRAIN A: Cognitive Brain (WHAT?) ──
        spec = CanonicalMissionSpec.compile_from_text(goal_text, mission_id=mission_id)
        
        # ── BRAIN B: Capability Brain (HAVE?) ──
        cap_eval = capability_brain.evaluate_mission_executability(spec)
        if not cap_eval.is_executable:
            return {
                "status": "BLOCKED",
                "phase": "CAPABILITY_AWARENESS",
                "reason": cap_eval.block_reason,
                "capability_gaps": cap_eval.capability_gaps,
                "executability_score": cap_eval.executability_score,
                "artifact_path": None
            }

        # ── BRAIN C: Execution Brain (HOW?) ──
        # 1. Adapt Schema
        raw_params = {
            "action": "create_file",
            "format": spec.artifact_type.lower(),
            "title": spec.objective,
            "charts": spec.chart_types,
            "add_charts": spec.chart_required
        }
        adapted = schema_adapter.adapt_office_intent(raw_params)

        # 2. Fast Governor Authorization
        intent = ActionIntent(
            intent_id=f"intent_{spec.spec_hash}",
            action_name="OFFICE_SUITE_MASTER",
            target_domain=TaskDomain.OFFICE,
            parameters=adapted,
            mutation=True
        )
        grant = fast_governor.evaluate_intent(intent)
        if not grant.allowed:
            return {
                "status": "DENIED",
                "phase": "GOVERNANCE",
                "reason": grant.reason,
                "artifact_path": None
            }

        # 3. Tool Execution via Smart Office Adapter
        out_file = None
        if spec.artifact_type == "XLSX":
            out_file = smart_office_adapter.synthesize_excel(
                title="Bảng Quản Lý Công Việc & Tiến Độ",
                filename=f"Quan_Ly_Tien_Do_{spec.mission_id}.xlsx",
                output_dir=output_dir,
                add_charts=spec.chart_required
            )

        # ── VERIFICATION (PROVE): Success Criteria Check ──
        criteria_results = []
        if out_file and os.path.exists(out_file):
            import openpyxl
            wb = openpyxl.load_workbook(out_file, data_only=False)
            ws = wb.active

            for crit in spec.success_criteria:
                crit.evaluated = True
                if crit.target_attribute == "file_exists":
                    crit.passed = os.path.exists(out_file) and os.path.getsize(out_file) > 0
                elif crit.target_attribute == "min_rows":
                    crit.passed = (ws.max_row >= (crit.expected_value + 3))
                elif crit.target_attribute == "has_formulas":
                    has_f = any(isinstance(c, str) and c.startswith("=") for row in ws.iter_rows(values_only=True) for c in row)
                    crit.passed = has_f
                elif crit.target_attribute == "has_charts":
                    crit.passed = (len(ws._charts) > 0)
                
                criteria_results.append({
                    "id": crit.criterion_id,
                    "desc": crit.description,
                    "passed": crit.passed
                })

        all_criteria_passed = all(c["passed"] for c in criteria_results) and (len(criteria_results) > 0)

        return {
            "status": "SUCCESS" if all_criteria_passed else "FAILED_VERIFICATION",
            "phase": "COMPLETED",
            "spec_hash": spec.spec_hash,
            "artifact_path": out_file,
            "matched_capabilities": cap_eval.matched_capabilities,
            "all_criteria_passed": all_criteria_passed,
            "criteria_results": criteria_results
        }


three_brain_orchestrator = ThreeBrainOrchestrator()
