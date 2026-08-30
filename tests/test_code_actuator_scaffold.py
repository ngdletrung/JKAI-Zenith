# -*- coding: utf-8 -*-
"""
Unit tests for Code-as-Action Sandbox and Artifact Verification.
"""

import os
import pytest
from core.kernel.code_actuator import code_actuator
from core.kernel.artifact_gate import verify_artifact_on_disk, format_delivery_receipt


def test_extract_python_code_markdown_block():
    raw = """Dưới đây là mã Python để tạo file:
```python
import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Test'
wb.save('test_extract.xlsx')
```
Chúc bạn thành công!"""
    code = code_actuator.extract_python_code(raw)
    assert code is not None
    assert "import openpyxl" in code
    assert "test_extract.xlsx" in code


def test_execute_office_code_excel_with_chart():
    excel_code = """
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Tiến Độ"

headers = ["STT", "Họ và tên", "Nhiệm vụ", "Tiến độ (%)", "Trạng thái"]
ws.append(headers)

# Style Header
header_fill = PatternFill(start_color="1B3A57", end_color="1B3A57", fill_type="solid")
header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Rows
data = [
    [1, "Nguyễn Văn A", "Backend API", 100, "Hoàn thành"],
    [2, "Trần Thị B", "Frontend UI", 90, "Đang xử lý"],
    [3, "Lê Văn C", "Database Design", 85, "Đang xử lý"]
]
for r in data:
    ws.append(r)

# Chart
chart = BarChart()
chart.title = "Tiến Độ Dự Án"
chart.x_axis.title = "Nhân sự"
chart.y_axis.title = "Tiến độ (%)"
data_ref = Reference(ws, min_col=4, min_row=1, max_row=4)
cats = Reference(ws, min_col=2, min_row=2, max_row=4)
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats)
ws.add_chart(chart, "G2")

wb.save("Test_Bao_Cao_Tien_Do.xlsx")
"""
    res = code_actuator.execute_office_code(excel_code, "Test_Bao_Cao_Tien_Do.xlsx")
    assert res["success"] is True
    assert res["file_path"] is not None
    assert os.path.exists(res["file_path"])
    assert res["filename"] == "Test_Bao_Cao_Tien_Do.xlsx"
    assert res["artifact_info"]["verified"] is True
    assert res["artifact_info"]["size_bytes"] > 0
    assert "/outputs/Test_Bao_Cao_Tien_Do.xlsx" in res["delivery_receipt"]


def test_execute_office_code_word_docx():
    word_code = """
import docx
doc = docx.Document()
doc.add_heading("BÁO CÁO TÁC CHIẾN JKAI", level=0)
doc.add_paragraph("Hệ điều hành JKAI đã sẵn sàng thực thi nhiệm vụ.")
doc.save("Test_Bao_Cao_Word.docx")
"""
    res = code_actuator.execute_office_code(word_code, "Test_Bao_Cao_Word.docx")
    assert res["success"] is True
    assert res["file_path"] is not None
    assert os.path.exists(res["file_path"])
    assert res["filename"] == "Test_Bao_Cao_Word.docx"
    assert res["artifact_info"]["verified"] is True
    assert res["artifact_info"]["size_bytes"] > 0
